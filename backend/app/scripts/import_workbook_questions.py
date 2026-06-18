from __future__ import annotations

import argparse
import json
import re
import sqlite3
import uuid
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional

from openpyxl import load_workbook

from backend.app.config import get_db_path
from backend.app.database import init_db, now_iso, to_json


WORKBOOK_GLOB = "AI素养题库设计*.xlsx"
SOURCE_LABEL = "AI素养题库设计.xlsx"

DIMENSION_BY_SHEET = {
    "理解AI": "理解AI（Understand）",
    "使用AI": "应用AI（Apply）",
    "创造AI": "创造AI（Create）",
    "AI伦理": "AI伦理（Ethics）",
}

DEFAULT_SECONDARY = {
    "理解AI（Understand）": "知道人工智能的基本概念（如机器学习、生成式AI等）",
    "应用AI（Apply）": "能判断某一任务是否适合使用AI完成",
    "创造AI（Create）": "能将真实问题转化为可由AI支持的问题形式",
    "AI伦理（Ethics）": "关注使用AI过程中的数据隐私与信息安全问题",
}

DIMENSION_PREFIX = {
    "理解AI（Understand）": "U",
    "应用AI（Apply）": "A",
    "创造AI（Create）": "C",
    "AI伦理（Ethics）": "E",
}

SECONDARY_BY_CODE = {
    "1.1": "知道人工智能的基本概念（如机器学习、生成式AI等）",
    "1.2": "理解AI基于数据进行训练和生成结果的基本机制",
    "1.3": "知道AI输出具有不确定性，可能出现错误或“幻觉”",
    "1.4": "能区分AI擅长与不擅长处理的任务类型",
    "2.1": "能判断某一任务是否适合使用AI完成",
    "2.2": "能清晰描述问题并与AI进行有效交互",
    "2.3": "能通过调整提问或指令优化AI输出结果",
    "2.4": "能对AI生成结果的准确性和合理性进行判断",
    "3.1": "能将真实问题转化为可由AI支持的问题形式",
    "3.2": "能利用AI生成并比较多种解决方案",
    "3.3": "能整合AI工具设计解决问题的流程或路径",
    "3.4": "能基于AI对方案进行调整、改进并形成较完整的解决方案",
    "4.1": "关注使用AI过程中的数据隐私与信息安全问题",
    "4.2": "能识别AI可能带来的偏见或不公平现象",
    "4.3": "在使用AI时遵循学术规范与诚信原则",
    "4.4": "意识到AI使用中的责任归属与潜在风险",
}


@dataclass
class ImportStats:
    parsed: int = 0
    skipped: int = 0
    inserted: int = 0
    updated: int = 0


def _clean(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).replace("\r\n", "\n").replace("\r", "\n")
    return re.sub(r"[ \t]+", " ", text).strip()


def _find_workbook(path_arg: Optional[str]) -> Path:
    if path_arg:
        path = Path(path_arg)
        if not path.exists():
            raise FileNotFoundError(path)
        return path

    candidates = sorted(Path("knowledge").glob(WORKBOOK_GLOB), key=lambda item: item.stat().st_size, reverse=True)
    if not candidates:
        raise FileNotFoundError(f"knowledge/{WORKBOOK_GLOB}")
    return candidates[0]


def _header_map(headers: Iterable[Any]) -> Dict[str, int]:
    return {_clean(header): index for index, header in enumerate(headers) if _clean(header)}


def _cell(row: List[Any], headers: Dict[str, int], *names: str) -> str:
    for name in names:
        index = headers.get(name)
        if index is not None and index < len(row):
            value = _clean(row[index])
            if value:
                return value
    return ""


def _parse_options(text: str) -> List[Dict[str, str]]:
    if not text:
        return []
    normalized = re.sub(r"\s+", " ", text).strip()
    matches = list(re.finditer(r"([A-Z])\s*[\.．、]\s*", normalized))
    if len(matches) < 2:
        return []

    options: List[Dict[str, str]] = []
    for index, match in enumerate(matches):
        start = match.end()
        end = matches[index + 1].start() if index + 1 < len(matches) else len(normalized)
        option_text = normalized[start:end].strip()
        if option_text:
            options.append({"id": match.group(1), "text": option_text})
    return options


def _difficulty(value: str) -> str:
    stars = value.count("★")
    if stars <= 1:
        return "easy"
    if stars == 2:
        return "medium"
    if stars >= 3:
        return "hard"
    lowered = value.lower()
    if lowered in {"easy", "medium", "hard"}:
        return lowered
    if "难" in value or "进阶" in value:
        return "hard"
    if "简" in value or "基础" in value:
        return "easy"
    return "medium"


def _stage_for_sheet(sheet: str) -> str:
    return "初中"


def load_knowledge_points() -> List[Dict[str, Any]]:
    with sqlite3.connect(str(get_db_path())) as connection:
        connection.row_factory = sqlite3.Row
        rows = connection.execute(
            """
            SELECT knowledge_code, stage, primary_dimension, secondary_dimension,
                   tertiary_ability, knowledge_point, description, source_row
            FROM knowledge_points
            ORDER BY stage ASC, primary_dimension ASC, source_row ASC
            """
        ).fetchall()
        return [dict(row) for row in rows]


def _match_knowledge_point(
    points: List[Dict[str, Any]],
    stage: str,
    dimension: str,
    knowledge_raw: str,
    context: str,
) -> Optional[Dict[str, Any]]:
    stage_points = [
        point
        for point in points
        if point["stage"] == stage and point["primary_dimension"] == dimension
    ]
    code_match = re.search(r"[1-4](?:\.[1-9]\d*)+", knowledge_raw)
    if code_match:
        stage_prefix = {"初中": "J", "高中": "H", "通用": "G"}.get(stage, "G")
        prefix = f"{stage_prefix}-{DIMENSION_PREFIX.get(dimension, 'X')}-{code_match.group(0)}"
        candidates = [point for point in stage_points if point["knowledge_code"].startswith(prefix)]
        if candidates:
            return candidates[0]

    raw = knowledge_raw.strip()
    if raw:
        for field in ["knowledge_point", "tertiary_ability", "secondary_dimension"]:
            candidates = [point for point in stage_points if point[field] == raw]
            if candidates:
                return candidates[0]
        candidates = [
            point
            for point in stage_points
            if raw in point["secondary_dimension"]
            or raw in point["tertiary_ability"]
            or raw in point["knowledge_point"]
            or point["knowledge_point"] in raw
        ]
        if candidates:
            return candidates[0]

    for point in stage_points:
        if point["knowledge_point"] and point["knowledge_point"] in context:
            return point
    return None


def _secondary(dimension: str, knowledge_raw: str, text: str) -> str:
    code_match = re.search(r"[1-4]\.[1-4]", knowledge_raw)
    if code_match:
        return SECONDARY_BY_CODE.get(code_match.group(0), DEFAULT_SECONDARY[dimension])

    if knowledge_raw and not re.fullmatch(r"[\d.\s]+", knowledge_raw):
        return knowledge_raw[:180]

    if dimension == "理解AI（Understand）":
        if re.search(r"数据|训练|生成|算法|算力", text):
            return SECONDARY_BY_CODE["1.2"]
        if re.search(r"错误|幻觉|不确定|核验|验证", text):
            return SECONDARY_BY_CODE["1.3"]
        if re.search(r"适合|不适合|擅长|局限", text):
            return SECONDARY_BY_CODE["1.4"]
    if dimension == "应用AI（Apply）":
        if re.search(r"提示|提问|指令|追问|优化", text):
            return SECONDARY_BY_CODE["2.3"]
        if re.search(r"准确|合理|核验|判断|结果", text):
            return SECONDARY_BY_CODE["2.4"]
        if re.search(r"描述问题|交互", text):
            return SECONDARY_BY_CODE["2.2"]
    if dimension == "创造AI（Create）":
        if re.search(r"多种|比较|方案", text):
            return SECONDARY_BY_CODE["3.2"]
        if re.search(r"流程|路径|系统", text):
            return SECONDARY_BY_CODE["3.3"]
        if re.search(r"调整|改进|优化", text):
            return SECONDARY_BY_CODE["3.4"]
    if dimension == "AI伦理（Ethics）":
        if re.search(r"偏见|公平|歧视", text):
            return SECONDARY_BY_CODE["4.2"]
        if re.search(r"学术|诚信|版权|抄袭", text):
            return SECONDARY_BY_CODE["4.3"]
        if re.search(r"责任|风险|事故", text):
            return SECONDARY_BY_CODE["4.4"]
    return DEFAULT_SECONDARY[dimension]


def _title_from_text(sheet: str, row_number: int, question_type: str, matched_point: Optional[Dict[str, Any]], text: str) -> str:
    if matched_point and matched_point.get("knowledge_point"):
        return f"{sheet}-{question_type}-{row_number - 1:03d}-{matched_point['knowledge_point'][:16]}"
    compact = re.sub(r"\s+", "", text)
    if compact:
        return f"{sheet}-{question_type}-{row_number - 1:03d}"
    return f"{sheet}第{row_number}题"


def _stable_id(sheet: str, row_number: int, stem: str) -> str:
    return str(uuid.uuid5(uuid.NAMESPACE_URL, f"{SOURCE_LABEL}:{sheet}:{row_number}:{stem[:120]}"))


def _question_from_row(
    sheet: str,
    row_number: int,
    row: List[Any],
    headers: Dict[str, int],
    last_type: str,
    knowledge_points: List[Dict[str, Any]],
) -> Optional[Dict[str, Any]]:
    dimension = DIMENSION_BY_SHEET.get(sheet)
    if not dimension:
        return None

    question_type = _cell(row, headers, "题型") or last_type or "未分类题"
    stem = _cell(row, headers, "题干")
    task_or_options = _cell(row, headers, "题项", "题项（选择题短题项后空四格）")
    answer = _cell(row, headers, "参考答案")
    scoring = _cell(row, headers, "评分标准")
    knowledge_raw = _cell(row, headers, "对应知识点")
    source = _cell(row, headers, "题目来源")
    difficulty_raw = _cell(row, headers, "难度系数")
    note = _cell(row, headers, "备注")

    if not stem:
        return None

    options = _parse_options(task_or_options)
    has_completion_signal = bool(answer or scoring or options or task_or_options)
    if not has_completion_signal:
        return None

    if options:
        question_text = stem
        scenario = ""
    else:
        question_text = task_or_options or stem
        scenario = stem if task_or_options else ""

    explanation_parts = []
    if answer:
        explanation_parts.append(f"参考答案：{answer}")
    if scoring:
        explanation_parts.append(f"评分标准：{scoring}")
    if note:
        explanation_parts.append(f"备注：{note}")
    explanation = "\n\n".join(explanation_parts) or "来源表未提供解析，需人工复核。"

    context = " ".join([stem, task_or_options, answer, scoring, knowledge_raw])
    stage = _stage_for_sheet(sheet)
    matched_point = _match_knowledge_point(knowledge_points, stage, dimension, knowledge_raw, context)
    secondary = matched_point["secondary_dimension"] if matched_point else _secondary(dimension, knowledge_raw, context)
    tertiary_ability = matched_point["tertiary_ability"] if matched_point else ""
    knowledge_point = matched_point["knowledge_point"] if matched_point else knowledge_raw
    source_reference = f"{SOURCE_LABEL} / {sheet} / 第{row_number}行"
    if source:
        source_reference = f"{source_reference}；{source}"

    return {
        "id": _stable_id(sheet, row_number, stem),
        "questionType": question_type,
        "title": _title_from_text(sheet, row_number, question_type, matched_point, stem),
        "question": question_text,
        "scenario": scenario,
        "options": options,
        "correctAnswer": answer,
        "explanation": explanation,
        "dimension": dimension,
        "secondaryDimension": secondary,
        "tertiaryDimension": tertiary_ability,
        "quaternaryDimension": knowledge_point,
        "subSkill": tertiary_ability or secondary,
        "cognitiveLevel": "apply" if sheet != "理解AI" else "understand",
        "difficultyEstimate": _difficulty(difficulty_raw),
        "tags": [sheet, "Excel导入", question_type],
        "knowledgePoints": [item for item in [knowledge_point, knowledge_raw] if item],
        "sourceReference": source_reference,
        "status": "reviewed",
        "stage": stage,
        "knowledgeCode": matched_point["knowledge_code"] if matched_point else "",
        "primaryDimension": dimension,
        "tertiaryAbility": tertiary_ability,
        "knowledgePoint": knowledge_point,
        "questionTask": task_or_options if not options else "",
        "referenceAnswer": answer,
        "scoringCriteria": scoring,
        "form": "选择题" if options else "文本",
        "scoreMax": None,
        "discrimination": None,
        "sourceSheet": sheet,
        "sourceRow": row_number,
    }


def parse_questions(workbook_path: Path) -> List[Dict[str, Any]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    knowledge_points = load_knowledge_points()
    questions: List[Dict[str, Any]] = []
    for sheet in DIMENSION_BY_SHEET:
        if sheet not in workbook.sheetnames:
            continue
        ws = workbook[sheet]
        headers = _header_map(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))
        last_type = ""
        for row_number, row_values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            row = list(row_values)
            explicit_type = _cell(row, headers, "题型")
            if explicit_type:
                last_type = explicit_type
            item = _question_from_row(sheet, row_number, row, headers, last_type, knowledge_points)
            if item:
                questions.append(item)
    return questions


def _next_item_code(connection: sqlite3.Connection) -> str:
    rows = connection.execute("SELECT item_code FROM questions").fetchall()
    max_code = 0
    for row in rows:
        value = row["item_code"]
        if value and value.isdigit():
            max_code = max(max_code, int(value))
    return str(max_code + 1).zfill(5)


def import_questions(questions: List[Dict[str, Any]], dry_run: bool = False) -> ImportStats:
    stats = ImportStats(parsed=len(questions))
    now = now_iso()
    with sqlite3.connect(str(get_db_path())) as connection:
        connection.row_factory = sqlite3.Row
        for question in questions:
            existing = connection.execute("SELECT item_code FROM questions WHERE id = ?", (question["id"],)).fetchone()
            item_code = existing["item_code"] if existing else _next_item_code(connection)
            if existing:
                stats.updated += 1
            else:
                stats.inserted += 1

            if dry_run:
                continue

            columns = [
                "id",
                "item_code",
                "question_type",
                "title",
                "question",
                "scenario",
                "options_json",
                "correct_answer",
                "explanation",
                "dimension",
                "secondary_dimension",
                "sub_skill",
                "tertiary_dimension",
                "quaternary_dimension",
                "cognitive_level",
                "difficulty_estimate",
                "tags_json",
                "knowledge_points_json",
                "source_reference",
                "status",
                "stage",
                "knowledge_code",
                "primary_dimension",
                "tertiary_ability",
                "knowledge_point",
                "question_task",
                "reference_answer",
                "scoring_criteria",
                "form",
                "score_max",
                "discrimination",
                "source_sheet",
                "source_row",
                "created_at",
                "updated_at",
            ]
            values = [
                question["id"],
                item_code,
                question["questionType"],
                question["title"],
                question["question"],
                question["scenario"],
                to_json(question["options"]),
                question["correctAnswer"],
                question["explanation"],
                question["dimension"],
                question["secondaryDimension"],
                question["subSkill"],
                question["tertiaryDimension"],
                question["quaternaryDimension"],
                question["cognitiveLevel"],
                question["difficultyEstimate"],
                to_json(question["tags"]),
                to_json(question["knowledgePoints"]),
                question["sourceReference"],
                question["status"],
                question["stage"],
                question["knowledgeCode"],
                question["primaryDimension"],
                question["tertiaryAbility"],
                question["knowledgePoint"],
                question["questionTask"],
                question["referenceAnswer"],
                question["scoringCriteria"],
                question["form"],
                question["scoreMax"],
                question["discrimination"],
                question["sourceSheet"],
                question["sourceRow"],
                now,
                now,
            ]
            placeholders = ", ".join("?" for _ in columns)
            connection.execute(
                f"INSERT OR REPLACE INTO questions ({', '.join(columns)}) VALUES ({placeholders})",
                values,
            )
    return stats


def main() -> int:
    parser = argparse.ArgumentParser(description="Import completed questions from AI literacy workbook into SQLite question bank.")
    parser.add_argument("--workbook", help="Path to AI素养题库设计.xlsx. Defaults to knowledge/AI素养题库设计*.xlsx")
    parser.add_argument("--dry-run", action="store_true", help="Parse and report without writing to the database.")
    parser.add_argument("--json-output", default="data/imported_workbook_questions.json", help="Write parsed question payloads for audit.")
    args = parser.parse_args()

    init_db()
    workbook_path = _find_workbook(args.workbook)
    questions = parse_questions(workbook_path)
    Path(args.json_output).parent.mkdir(parents=True, exist_ok=True)
    Path(args.json_output).write_text(json.dumps(questions, ensure_ascii=False, indent=2), encoding="utf-8")
    stats = import_questions(questions, dry_run=args.dry_run)

    by_dimension: Dict[str, int] = {}
    by_type: Dict[str, int] = {}
    for question in questions:
        by_dimension[question["dimension"]] = by_dimension.get(question["dimension"], 0) + 1
        by_type[question["questionType"]] = by_type.get(question["questionType"], 0) + 1

    print(f"workbook={workbook_path}")
    print(f"dry_run={args.dry_run}")
    print(f"parsed={stats.parsed}")
    print(f"inserted={stats.inserted}")
    print(f"updated={stats.updated}")
    print(f"json_output={args.json_output}")
    print("by_dimension=")
    for key, value in sorted(by_dimension.items()):
        print(f"  {key}: {value}")
    print("by_type=")
    for key, value in sorted(by_type.items()):
        print(f"  {key}: {value}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
