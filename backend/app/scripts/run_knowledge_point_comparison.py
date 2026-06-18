from __future__ import annotations

import argparse
import asyncio
import json
import re
import shutil
import urllib.error
import urllib.request
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from backend.app.config import (
    get_deepseek_api_key,
    get_deepseek_base_url,
    get_deepseek_model,
    get_xfyun_api_key,
    get_xfyun_base_url,
    get_xfyun_model,
)
from backend.app.database import init_db
from backend.app.repositories import list_knowledge_points


DATE_STAMP = datetime.now().strftime("%Y-%m-%d")
PLAN_PATH = Path(f"data/knowledge_point_comparison_plan_{DATE_STAMP}.json")
ROWS_PATH = Path(f"data/knowledge_point_model_comparison_rows_{DATE_STAMP}.json")
OUTPUT_PATH = Path(f"outputs/knowledge_point_model_comparison_{DATE_STAMP}.xlsx")
DISPLAY_COPY_PATH = Path(f"outputs/知识点总表出题模型对比_{DATE_STAMP}.xlsx")

DIMENSIONS = ["理解AI（Understand）", "应用AI（Apply）", "创造AI（Create）", "AI伦理（Ethics）"]
DIM_CODE = {
    "理解AI（Understand）": "U",
    "应用AI（Apply）": "A",
    "创造AI（Create）": "C",
    "AI伦理（Ethics）": "E",
}
PROVIDER_LABELS = {"xfyun": "讯飞", "deepseek": "DeepSeek", "codex": "Codex"}


def _strip_code_fence(content: str) -> str:
    cleaned = content.strip()
    cleaned = re.sub(r"^```(?:json)?\s*", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"\s*```$", "", cleaned)
    return cleaned.strip()


def _extract_json(content: str) -> Dict[str, Any]:
    cleaned = _strip_code_fence(content)
    try:
        return json.loads(cleaned)
    except json.JSONDecodeError:
        start = cleaned.find("{")
        end = cleaned.rfind("}")
        if start >= 0 and end > start:
            return json.loads(cleaned[start : end + 1])
        raise


def _normalize_options(value: Any) -> List[Dict[str, str]]:
    if isinstance(value, dict):
        return [{"id": str(key).strip(), "text": str(text).strip()} for key, text in value.items()]
    if isinstance(value, list):
        options = []
        for index, item in enumerate(value):
            default_id = chr(65 + index)
            if isinstance(item, dict):
                option_id = str(item.get("id") or item.get("key") or item.get("label") or default_id).strip()
                text = str(item.get("text") or item.get("content") or item.get("value") or "").strip()
            else:
                option_id = default_id
                text = str(item).strip()
            if option_id and text:
                options.append({"id": option_id, "text": text})
        return options
    return []


def _normalize_answer(value: Any) -> str:
    if isinstance(value, list):
        return "".join(str(item).strip() for item in value if str(item).strip())
    return str(value or "").strip()


def _point_to_dict(point) -> Dict[str, Any]:
    return {
        "id": point.id,
        "knowledgeCode": point.knowledgeCode,
        "stage": point.stage,
        "primaryDimension": point.primaryDimension,
        "secondaryDimension": point.secondaryDimension,
        "tertiaryAbility": point.tertiaryAbility,
        "knowledgePoint": point.knowledgePoint,
        "description": point.description,
        "cognitiveLevel": point.cognitiveLevel,
        "suggestedQuestionTypes": point.suggestedQuestionTypes,
        "sourceReferences": point.sourceReferences,
        "sourceSheet": point.sourceSheet,
        "sourceRow": point.sourceRow,
    }


def select_points(per_dimension: int = 5, stage: str = "初中") -> List[Dict[str, Any]]:
    init_db()
    points = [_point_to_dict(point) for point in list_knowledge_points(stage=stage)]
    selected: List[Dict[str, Any]] = []
    for dimension in DIMENSIONS:
        candidates = [point for point in points if point["primaryDimension"] == dimension]
        candidates.sort(key=lambda item: (item["sourceRow"], item["knowledgeCode"]))
        selected.extend(candidates[:per_dimension])
    return selected


def build_prompt(points: List[Dict[str, Any]]) -> List[Dict[str, str]]:
    point_lines = []
    for index, point in enumerate(points, start=1):
        point_lines.append(
            "\n".join(
                [
                    f"{index}. knowledgeCode: {point['knowledgeCode']}",
                    f"一级维度: {point['primaryDimension']}",
                    f"二级维度: {point['secondaryDimension']}",
                    f"三级能力: {point['tertiaryAbility']}",
                    f"四级知识点: {point['knowledgePoint']}",
                    f"知识点说明: {point['description']}",
                    f"来源: {'；'.join(point['sourceReferences'])}",
                ]
            )
        )

    user_prompt = f"""
请基于下面 5 个 AI 素养知识点，为每个 knowledgeCode 各生成 1 道初中学生适用的场景化单选题，共 5 道。

硬性要求：
1. 每道题必须严格对应一个给定 knowledgeCode，不要遗漏、不要新增 knowledgeCode。
2. 题目必须是单选题，options 必须是 A/B/C/D 四个选项。
3. correctAnswer 只能是 A、B、C、D 之一。
4. 场景要贴近中学生学习、校园生活或日常 AI 使用。
5. 不要直接照抄知识点说明，要把知识点转化为判断、应用或辨析情境。
6. dimension、secondaryDimension、tertiaryDimension、quaternaryDimension 必须沿用给定知识点。
7. 只输出合法 JSON，不要 Markdown，不要解释。

输出 JSON 格式：
{{
  "questions": [
    {{
      "knowledgeCode": "...",
      "questionType": "单选",
      "title": "短标题",
      "scenario": "场景",
      "question": "题干问题",
      "options": [{{"id":"A","text":"..."}}, {{"id":"B","text":"..."}}, {{"id":"C","text":"..."}}, {{"id":"D","text":"..."}}],
      "correctAnswer": "A",
      "explanation": "解析",
      "dimension": "一级维度",
      "secondaryDimension": "二级维度",
      "tertiaryDimension": "三级能力",
      "quaternaryDimension": "四级知识点",
      "knowledgePoints": ["四级知识点"],
      "difficultyEstimate": "easy/medium/hard"
    }}
  ]
}}

知识点：
{"\n\n".join(point_lines)}
""".strip()

    return [
        {
            "role": "system",
            "content": "你是 AI 素养测评题库出题专家。你只输出合法 JSON object。",
        },
        {"role": "user", "content": user_prompt},
    ]


async def call_chat_completion(provider: str, points: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    if provider == "xfyun":
        api_key = get_xfyun_api_key()
        base_url = get_xfyun_base_url()
        model = get_xfyun_model()
        payload = {
            "model": model,
            "messages": build_prompt(points),
            "temperature": 0.35,
        }
    elif provider == "deepseek":
        api_key = get_deepseek_api_key()
        base_url = get_deepseek_base_url()
        model = get_deepseek_model()
        payload = {
            "model": model,
            "messages": build_prompt(points),
            "temperature": 0.35,
            "response_format": {"type": "json_object"},
        }
    else:
        raise ValueError(f"Unsupported provider: {provider}")

    result = await asyncio.to_thread(
        _post_json,
        f"{base_url}/chat/completions",
        {"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"},
        payload,
    )
    content = result.get("choices", [{}])[0].get("message", {}).get("content", "")
    parsed = _extract_json(content)
    questions = parsed.get("questions")
    if not isinstance(questions, list):
        raise RuntimeError(f"{provider} response missing questions array")
    return [normalize_generated_question(provider, question, points) for question in questions]


def _post_json(url: str, headers: Dict[str, str], payload: Dict[str, Any]) -> Dict[str, Any]:
    request = urllib.request.Request(
        url,
        data=json.dumps(payload, ensure_ascii=False).encode("utf-8"),
        headers=headers,
        method="POST",
    )
    try:
        with urllib.request.urlopen(request, timeout=180) as response:
            body = response.read().decode("utf-8")
            return json.loads(body)
    except urllib.error.HTTPError as exc:
        body = exc.read().decode("utf-8", errors="replace")
        try:
            payload = json.loads(body)
        except json.JSONDecodeError:
            payload = {}
        message = payload.get("error", {}).get("message") if isinstance(payload, dict) else ""
        raise RuntimeError(message or f"HTTP {exc.code}: {body[:300]}") from exc


def _option_set(correct_text: str, distractors: List[str], correct_id: str = "B") -> List[Dict[str, str]]:
    ids = ["A", "B", "C", "D"]
    texts = distractors[:]
    texts.insert(ids.index(correct_id), correct_text)
    return [{"id": option_id, "text": text} for option_id, text in zip(ids, texts)]


def codex_generate(points: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    questions = []
    for point in points:
        dimension = point["primaryDimension"]
        kp = point["knowledgePoint"]
        tertiary = point["tertiaryAbility"]
        description = point["description"] or kp
        if dimension == "理解AI（Understand）":
            scenario = f"信息科技课上，老师让同学们解释“{kp}”与日常使用 AI 工具之间的关系。"
            question = f"下列哪种说法最符合“{kp}”这一知识点？"
            correct = f"能结合“{description}”说明 AI 的基本概念或运行机制，而不是只记住工具名称。"
            distractors = [
                "只要某个软件看起来很聪明，就一定具有人类一样的理解能力。",
                "判断 AI 时只需要看它回答得快不快，不需要了解背后的数据或算法。",
                "所有自动化程序都等同于人工智能，二者没有区别。",
            ]
        elif dimension == "应用AI（Apply）":
            scenario = f"小组作业中，学生准备用 AI 辅助完成一项学习任务，任务涉及“{kp}”。"
            question = "下列哪种做法最能体现合理应用 AI？"
            correct = f"先明确任务目标和限制，再根据“{description}”判断 AI 能辅助哪一部分，并保留人工核验。"
            distractors = [
                "把整个作业直接交给 AI 完成，不再检查输出内容。",
                "只要 AI 给出答案，就默认它比同学和老师都更可靠。",
                "为了省事，向 AI 输入同学的隐私信息来换取更具体的结果。",
            ]
        elif dimension == "创造AI（Create）":
            scenario = f"学校希望学生围绕“{kp}”设计一个 AI 辅助解决校园问题的小方案。"
            question = "哪一种设计思路最符合创造 AI 维度的要求？"
            correct = f"把真实问题拆解为 AI 可支持的任务，说明所需数据、流程、人工职责和改进方式。"
            distractors = [
                "先决定使用最热门的 AI 工具，再想办法把所有问题都塞给它。",
                "只写一个口号式方案，不说明数据、流程或评价方式。",
                "让 AI 自动替代所有师生决策，避免人工参与带来的麻烦。",
            ]
        else:
            scenario = f"班级讨论使用 AI 时的规范问题，其中涉及“{kp}”。"
            question = "下列哪种判断最符合 AI 伦理要求？"
            correct = f"在使用 AI 前识别隐私、公平、责任或诚信风险，并采取相应保护措施。"
            distractors = [
                "只要 AI 能提高效率，就可以忽略数据来源和使用边界。",
                "AI 输出的问题都应由系统负责，使用者不需要承担任何责任。",
                "为了得到更准确的结果，可以随意上传他人的个人信息。",
            ]

        item = {
            "knowledgeCode": point["knowledgeCode"],
            "questionType": "单选",
            "title": f"{kp}情境辨析",
            "scenario": scenario,
            "question": question,
            "options": _option_set(correct, distractors),
            "correctAnswer": "B",
            "explanation": f"本题考查“{kp}”。关键是把知识点放回具体情境中判断：{description}",
            "dimension": dimension,
            "secondaryDimension": point["secondaryDimension"],
            "tertiaryDimension": tertiary,
            "quaternaryDimension": kp,
            "knowledgePoints": [kp],
            "difficultyEstimate": "medium",
        }
        questions.append(normalize_generated_question("codex", item, points))
    return questions


def normalize_generated_question(provider: str, raw: Dict[str, Any], points: List[Dict[str, Any]]) -> Dict[str, Any]:
    knowledge_code = str(raw.get("knowledgeCode") or raw.get("knowledge_code") or "").strip()
    point = next((item for item in points if item["knowledgeCode"] == knowledge_code), None)
    if not point:
        fallback_index = len([item for item in points if item.get("_used_by_" + provider)]) % len(points)
        point = points[fallback_index]
        point["_used_by_" + provider] = True
        knowledge_code = point["knowledgeCode"]

    return {
        "provider": provider,
        "source": PROVIDER_LABELS[provider],
        "knowledgeCode": knowledge_code,
        "knowledgeEntry": f"{point['knowledgeCode']} {point['knowledgePoint']}",
        "stage": point["stage"],
        "primaryDimension": point["primaryDimension"],
        "secondaryDimension": point["secondaryDimension"],
        "tertiaryAbility": point["tertiaryAbility"],
        "knowledgePoint": point["knowledgePoint"],
        "questionType": str(raw.get("questionType") or "单选").strip(),
        "title": str(raw.get("title") or point["knowledgePoint"]).strip(),
        "scenario": str(raw.get("scenario") or "").strip(),
        "question": str(raw.get("question") or "").strip(),
        "options": _normalize_options(raw.get("options")),
        "correctAnswer": _normalize_answer(raw.get("correctAnswer")),
        "explanation": str(raw.get("explanation") or "").strip(),
        "difficultyEstimate": str(raw.get("difficultyEstimate") or "medium").strip(),
        "sourceReference": "；".join(point["sourceReferences"]),
        "description": point["description"],
        "raw": raw,
    }


def assess(row: Dict[str, Any]) -> tuple[str, str, str]:
    issues = []
    if row["questionType"] != "单选":
        issues.append("题型不是单选")
    if len(row["options"]) != 4:
        issues.append(f"选项数为{len(row['options'])}")
    option_ids = {option["id"] for option in row["options"]}
    if row["correctAnswer"] not in option_ids:
        issues.append("答案未匹配选项")
    if not row["scenario"] or len(row["scenario"]) < 12:
        issues.append("场景偏短")
    if not row["question"]:
        issues.append("题干为空")
    if not row["explanation"] or len(row["explanation"]) < 20:
        issues.append("解析偏短")
    combined = " ".join([row["title"], row["scenario"], row["question"], row["explanation"]])
    if row["knowledgePoint"] not in combined and row["tertiaryAbility"] not in combined:
        issues.append("考点显性不足")

    if not issues:
        return "较好：结构完整，考点和场景基本清楚。", "建议入库", "可进入人工复审。"
    if len(issues) <= 2 and "答案未匹配选项" not in issues and "选项数" not in "；".join(issues):
        return "一般：主体可用，但需人工修改。", "修改后入库", "；".join(issues)
    return "较弱：结构或考点定位存在明显问题。", "暂不入库", "；".join(issues)


def number_rows(rows: List[Dict[str, Any]]) -> None:
    provider_order = {"xfyun": 0, "deepseek": 1, "codex": 2}
    rows.sort(
        key=lambda row: (
            provider_order[row["provider"]],
            DIMENSIONS.index(row["primaryDimension"]),
            row["knowledgeCode"],
        )
    )
    counts = Counter()
    for row in rows:
        counts[(row["provider"], row["primaryDimension"])] += 1
        row["number"] = f"{row['provider'].upper()}-{DIM_CODE[row['primaryDimension']]}-{counts[(row['provider'], row['primaryDimension'])]:02d}"
        quality, recommend, note = assess(row)
        row["quality"] = quality
        row["recommend"] = recommend
        row["note"] = note
        row["optionsText"] = "\n".join(f"{option['id']}. {option['text']}" for option in row["options"])


async def generate_api_rows(points: List[Dict[str, Any]], provider: str) -> List[Dict[str, Any]]:
    rows: List[Dict[str, Any]] = []
    for dimension in DIMENSIONS:
        batch = [point for point in points if point["primaryDimension"] == dimension]
        print(f"[{provider}] generating {dimension}: {len(batch)}")
        rows.extend(await call_chat_completion(provider, batch))
    return rows


async def generate_rows(points: List[Dict[str, Any]], skip_api: bool = False) -> List[Dict[str, Any]]:
    rows: List[Dict[str, Any]] = []
    if not skip_api:
        rows.extend(await generate_api_rows(points, "xfyun"))
        rows.extend(await generate_api_rows(points, "deepseek"))
    rows.extend(codex_generate(points))
    number_rows(rows)
    return rows


DETAIL_HEADERS = [
    "编号",
    "来源",
    "knowledgeCode",
    "知识条目",
    "学段",
    "一级维度",
    "二级维度",
    "三级能力",
    "四级知识点",
    "题型",
    "题目",
    "场景",
    "选项",
    "正确答案",
    "解析",
    "参考知识点",
    "初步质量评价",
    "是否建议入库",
    "问题备注",
]

FIELD_MAP = {
    "编号": "number",
    "来源": "source",
    "knowledgeCode": "knowledgeCode",
    "知识条目": "knowledgeEntry",
    "学段": "stage",
    "一级维度": "primaryDimension",
    "二级维度": "secondaryDimension",
    "三级能力": "tertiaryAbility",
    "四级知识点": "knowledgePoint",
    "题型": "questionType",
    "题目": "question",
    "场景": "scenario",
    "选项": "optionsText",
    "正确答案": "correctAnswer",
    "解析": "explanation",
    "参考知识点": "description",
    "初步质量评价": "quality",
    "是否建议入库": "recommend",
    "问题备注": "note",
}


def style_header(row) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)
    for cell in row:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def style_sheet(ws) -> None:
    thin = Side(style="thin", color="D9E2F3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def build_excel(rows: List[Dict[str, Any]], output_path: Path = OUTPUT_PATH) -> None:
    wb = Workbook()
    detail = wb.active
    detail.title = "对比明细"
    detail.append(DETAIL_HEADERS)
    style_header(detail[1])
    for row in rows:
        detail.append([row.get(FIELD_MAP[header], "") for header in DETAIL_HEADERS])

    widths = [15, 12, 22, 28, 10, 20, 28, 28, 22, 14, 42, 42, 52, 12, 54, 42, 32, 16, 36]
    for index, width in enumerate(widths, start=1):
        detail.column_dimensions[get_column_letter(index)].width = width
    for row_index in range(2, detail.max_row + 1):
        detail.row_dimensions[row_index].height = 90
    style_sheet(detail)

    summary = wb.create_sheet("汇总")
    summary.append(["来源", "一级维度", "题目数", "建议入库", "修改后入库", "暂不入库"])
    style_header(summary[1])
    for source in ["讯飞", "DeepSeek", "Codex"]:
        for dimension in DIMENSIONS:
            filtered = [row for row in rows if row["source"] == source and row["primaryDimension"] == dimension]
            counts = Counter(row["recommend"] for row in filtered)
            summary.append(
                [
                    source,
                    dimension,
                    len(filtered),
                    counts.get("建议入库", 0),
                    counts.get("修改后入库", 0),
                    counts.get("暂不入库", 0),
                ]
            )
    summary.append([])
    summary.append(["总题数", len(rows)])
    summary.append(["说明", "三组均基于同一批 knowledge_points：每个 UACE 一级维度 5 个知识点，每个知识点每组生成 1 道题。"])
    for index, width in enumerate([14, 24, 12, 12, 12, 12], start=1):
        summary.column_dimensions[get_column_letter(index)].width = width
    style_sheet(summary)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    shutil.copyfile(output_path, DISPLAY_COPY_PATH)

    check = load_workbook(output_path, read_only=True, data_only=True)
    if check["对比明细"].max_row - 1 != len(rows):
        raise RuntimeError("Workbook row count verification failed")


def save_json(path: Path, data: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


async def main_async(args: argparse.Namespace) -> int:
    points = select_points(per_dimension=args.per_dimension, stage=args.stage)
    if len(points) != args.per_dimension * len(DIMENSIONS):
        raise RuntimeError(f"Expected {args.per_dimension * len(DIMENSIONS)} points, selected {len(points)}")
    save_json(PLAN_PATH, points)
    print(f"plan={PLAN_PATH}")
    print(f"selected_points={len(points)}")

    rows = await generate_rows(points, skip_api=args.skip_api)
    save_json(ROWS_PATH, rows)
    build_excel(rows)

    print(f"rows={len(rows)}")
    print(f"rows_json={ROWS_PATH}")
    print(f"excel={OUTPUT_PATH}")
    print(f"display_copy={DISPLAY_COPY_PATH}")
    print("source_counts=", Counter(row["source"] for row in rows))
    print("dimension_counts=")
    for key, value in sorted(Counter((row["source"], row["primaryDimension"]) for row in rows).items()):
        print(f"  {key}: {value}")
    return 0


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Run a three-provider question generation comparison from knowledge_points.")
    parser.add_argument("--stage", default="初中")
    parser.add_argument("--per-dimension", type=int, default=5)
    parser.add_argument("--skip-api", action="store_true", help="Only generate the local Codex/template group.")
    return parser


def main() -> int:
    args = build_parser().parse_args()
    return asyncio.run(main_async(args))


if __name__ == "__main__":
    raise SystemExit(main())
