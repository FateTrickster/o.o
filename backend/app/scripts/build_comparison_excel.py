from __future__ import annotations

import json
import shutil
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


INPUT = Path("data/model_comparison_rows_2026-06-12.json")
OUTPUT = Path("outputs/model_comparison_2026-06-12.xlsx")
DISPLAY_COPY = Path("outputs/AI素养出题模型对比_2026-06-12.xlsx")


DETAIL_HEADERS = [
    "编号",
    "来源",
    "计划维度",
    "题目自填维度",
    "知识条目",
    "题目",
    "场景",
    "选项",
    "正确答案",
    "解析",
    "参考知识点",
    "初步质量评价",
    "是否建议入库",
    "问题备注",
]


FIELD_MAP = {
    "编号": "number",
    "来源": "source",
    "计划维度": "intendedDimension",
    "题目自填维度": "dimension",
    "知识条目": "knowledgeEntry",
    "题目": "question",
    "场景": "scenario",
    "选项": "options",
    "正确答案": "correctAnswer",
    "解析": "explanation",
    "参考知识点": "knowledgePoints",
    "初步质量评价": "quality",
    "是否建议入库": "recommend",
    "问题备注": "note",
}


WIDTHS = {
    "编号": 15,
    "来源": 12,
    "计划维度": 20,
    "题目自填维度": 20,
    "知识条目": 24,
    "题目": 42,
    "场景": 44,
    "选项": 52,
    "正确答案": 12,
    "解析": 54,
    "参考知识点": 42,
    "初步质量评价": 34,
    "是否建议入库": 16,
    "问题备注": 40,
}


def _style_header(row):
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)
    alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for cell in row:
        cell.fill = fill
        cell.font = font
        cell.alignment = alignment


def _style_sheet(ws):
    thin = Side(style="thin", color="D9E2F3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _recommend_fill(value: str):
    if value == "建议入库":
        return PatternFill("solid", fgColor="E2F0D9")
    if value == "修改后入库":
        return PatternFill("solid", fgColor="FFF2CC")
    return PatternFill("solid", fgColor="FCE4D6")


def build_workbook(rows):
    wb = Workbook()
    detail = wb.active
    detail.title = "对比明细"
    detail.append(DETAIL_HEADERS)
    _style_header(detail[1])

    for item in rows:
        detail.append([item.get(FIELD_MAP[header], "") for header in DETAIL_HEADERS])

    for index, header in enumerate(DETAIL_HEADERS, start=1):
        detail.column_dimensions[get_column_letter(index)].width = WIDTHS[header]

    recommendation_col = DETAIL_HEADERS.index("是否建议入库") + 1
    for row_idx in range(2, detail.max_row + 1):
        detail.row_dimensions[row_idx].height = 88
        recommendation = detail.cell(row=row_idx, column=recommendation_col).value
        detail.cell(row=row_idx, column=recommendation_col).fill = _recommend_fill(recommendation)

    _style_sheet(detail)

    summary = wb.create_sheet("汇总")
    summary_headers = ["来源", "计划维度", "题目数", "建议入库", "修改后入库", "暂不入库"]
    summary.append(summary_headers)
    _style_header(summary[1])

    sources = ["讯飞", "DeepSeek", "Codex"]
    dimensions = ["理解AI（Understand）", "应用AI（Apply）", "创造AI（Create）", "AI伦理（Ethics）"]
    for source in sources:
        for dimension in dimensions:
            filtered = [row for row in rows if row["source"] == source and row["intendedDimension"] == dimension]
            counts = Counter(row["recommend"] for row in filtered)
            summary.append(
                [
                    source,
                    dimension,
                    len(filtered),
                    counts.get("建议入库", 0),
                    counts.get("修改后入库", 0),
                    counts.get("暂不入库", 0),
                ]
            )

    summary.append([])
    summary.append(["总题数", len(rows)])
    summary.append(["说明", "三组题目均基于同一个知识条目：AI素养框架.docx #1；每组每个 UACE 维度 5 道。"])

    for col_idx in range(1, 7):
        summary.column_dimensions[get_column_letter(col_idx)].width = [14, 24, 12, 12, 12, 12][col_idx - 1]
    _style_sheet(summary)
    return wb


def main() -> int:
    rows = json.loads(INPUT.read_text(encoding="utf-8"))
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb = build_workbook(rows)
    wb.save(OUTPUT)
    shutil.copyfile(OUTPUT, DISPLAY_COPY)

    check = load_workbook(OUTPUT, read_only=True)
    detail = check["对比明细"]
    summary = check["汇总"]
    detail_rows = detail.max_row - 1
    source_counts = Counter(row["source"] for row in rows)
    dimension_counts = Counter((row["source"], row["intendedDimension"]) for row in rows)

    print(f"output={OUTPUT}")
    print(f"display_copy={DISPLAY_COPY}")
    print(f"detail_rows={detail_rows}")
    print(f"summary_rows={summary.max_row}")
    print("source_counts=", source_counts)
    for key, value in sorted(dimension_counts.items()):
        print(key, value)
    if detail_rows != 60:
        raise RuntimeError(f"Expected 60 detail rows, got {detail_rows}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
