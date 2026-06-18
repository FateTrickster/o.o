from __future__ import annotations

import argparse
import json
import re
import sqlite3
import uuid
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional

from openpyxl import load_workbook

from backend.app.config import get_db_path
from backend.app.database import init_db, now_iso, to_json


DIMENSION_ALIASES = {
    "Understand": "理解AI（Understand）",
    "理解AI": "理解AI（Understand）",
    "1.理解AI": "理解AI（Understand）",
    "1 理解AI": "理解AI（Understand）",
    "Apply": "应用AI（Apply）",
    "应用AI": "应用AI（Apply）",
    "使用AI": "应用AI（Apply）",
    "2.使用AI": "应用AI（Apply）",
    "2 使用AI": "应用AI（Apply）",
    "Create": "创造AI（Create）",
    "创造AI": "创造AI（Create）",
    "3.创造AI": "创造AI（Create）",
    "3 创造AI": "创造AI（Create）",
    "Ethics": "AI伦理（Ethics）",
    "AI伦理": "AI伦理（Ethics）",
    "4.AI伦理": "AI伦理（Ethics）",
    "4 AI伦理": "AI伦理（Ethics）",
}

DIMENSION_PREFIX = {
    "理解AI（Understand）": "U",
    "应用AI（Apply）": "A",
    "创造AI（Create）": "C",
    "AI伦理（Ethics）": "E",
}


@dataclass
class ImportStats:
    parsed: int = 0
    inserted: int = 0
    updated: int = 0


def _clean(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).replace("\r", "\n")).strip()


def _split_list(value: str) -> List[str]:
    if not value:
        return []
    return [item.strip() for item in re.split(r"[；;]\s*", value) if item.strip()]


def _strip_leading_code(value: str) -> str:
    return re.sub(r"^\s*\d+(?:\.\d+)*\s*", "", value).strip()


def _extract_leading_code(value: str) -> str:
    match = re.match(r"\s*(\d+(?:\.\d+)*)", value or "")
    return match.group(1) if match else ""


def _normalize_dimension(value: str) -> str:
    cleaned = _strip_leading_code(value).replace(" ", "")
    if value in DIMENSION_ALIASES:
        return DIMENSION_ALIASES[value]
    if cleaned in DIMENSION_ALIASES:
        return DIMENSION_ALIASES[cleaned]
    if "伦理" in value or "Ethics" in value:
        return "AI伦理（Ethics）"
    if "创造" in value or "Create" in value:
        return "创造AI（Create）"
    if "应用" in value or "使用" in value or "Apply" in value:
        return "应用AI（Apply）"
    return "理解AI（Understand）"


def _header_map(headers: Iterable[Any]) -> Dict[str, int]:
    return {_clean(header): index for index, header in enumerate(headers) if _clean(header)}


def _cell(row: List[Any], headers: Dict[str, int], *names: str) -> str:
    for name in names:
        index = headers.get(name)
        if index is not None and index < len(row):
            value = _clean(row[index])
            if value:
                return value
    return ""


def _stable_id(source: str, source_sheet: str, source_row: int, knowledge_point: str) -> str:
    return str(uuid.uuid5(uuid.NAMESPACE_URL, f"{source}:{source_sheet}:{source_row}:{knowledge_point}"))


def _knowledge_code(stage: str, dimension: str, raw_code: str, sequence: int) -> str:
    prefix = DIMENSION_PREFIX.get(dimension, "X")
    stage_prefix = {"初中": "J", "高中": "H", "通用": "G"}.get(stage, "G")
    if raw_code:
        return f"{stage_prefix}-{prefix}-{raw_code}-{sequence:04d}"
    return f"{stage_prefix}-{prefix}-{sequence:04d}"


def _find_workbooks(root: Path) -> Dict[str, Path]:
    files = list(root.glob("*.xlsx"))
    result: Dict[str, Path] = {}
    for path in files:
        name = path.name
        if "知识图谱" in name:
            result["graph"] = path
        if "题库设计" in name:
            result["design"] = path
    return result


def _make_point(
    *,
    source: str,
    stage: str,
    primary_dimension: str,
    secondary_dimension: str,
    tertiary_ability: str,
    knowledge_point: str,
    description: str,
    cognitive_level: str,
    suggested_question_types: List[str],
    source_references: List[str],
    tags: List[str],
    source_sheet: str,
    source_row: int,
    sequence: int,
    raw_code: str = "",
) -> Dict[str, Any]:
    code = _knowledge_code(stage, primary_dimension, raw_code, sequence)
    return {
        "id": _stable_id(source, source_sheet, source_row, knowledge_point or description),
        "knowledgeCode": code,
        "stage": stage,
        "primaryDimension": primary_dimension,
        "secondaryDimension": secondary_dimension,
        "tertiaryAbility": tertiary_ability,
        "knowledgePoint": _strip_leading_code(knowledge_point) or tertiary_ability or secondary_dimension,
        "description": description,
        "cognitiveLevel": cognitive_level,
        "suggestedQuestionTypes": suggested_question_types,
        "sourceReferences": source_references,
        "tags": [tag for tag in tags if tag],
        "status": "active",
        "sourceSheet": source_sheet,
        "sourceRow": source_row,
    }


def parse_graph_workbook(path: Path) -> List[Dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    ws = workbook.worksheets[0]
    headers = _header_map(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))
    points: List[Dict[str, Any]] = []
    for sequence, (row_number, row_values) in enumerate(
        ((index, list(values)) for index, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2)),
        start=1,
    ):
        primary = _cell(row_values, headers, "一级维度")
        secondary = _cell(row_values, headers, "二级观察条目")
        tertiary = _cell(row_values, headers, "三级能力单元")
        knowledge = _cell(row_values, headers, "四级知识点")
        question_types = _cell(row_values, headers, "对应主观题题型")
        if not knowledge:
            continue
        dimension = _normalize_dimension(primary)
        points.append(
            _make_point(
                source=path.name,
                stage="通用",
                primary_dimension=dimension,
                secondary_dimension=secondary,
                tertiary_ability=tertiary,
                knowledge_point=knowledge,
                description=f"{secondary} / {tertiary} / {knowledge}",
                cognitive_level="",
                suggested_question_types=_split_list(question_types),
                source_references=[path.name],
                tags=["知识图谱", primary],
                source_sheet=ws.title,
                source_row=row_number,
                sequence=sequence,
            )
        )
    return points


def parse_design_workbook(path: Path) -> List[Dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    points: List[Dict[str, Any]] = []
    sequence = 1
    for sheet_name in ["初中", "高中"]:
        if sheet_name not in workbook.sheetnames:
            continue
        ws = workbook[sheet_name]
        headers = _header_map(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))
        current_primary = ""
        current_secondary = ""
        current_tertiary = ""
        for row_number, row_values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            row = list(row_values)
            primary = _cell(row, headers, "一级知识点", "一级维度") or current_primary
            secondary = _cell(row, headers, "二级知识点", "二级维度") or current_secondary
            tertiary = _cell(row, headers, "三级知识点", "三级维度") or current_tertiary
            knowledge = _cell(row, headers, "四级知识点", "具体知识点")
            description = _cell(row, headers, "具体描述", "知识点描述")
            source_reference = _cell(row, headers, "知识点来源")
            cognitive_level = _cell(row, headers, "认知等级")
            note = _cell(row, headers, "测评建议", "问题标注")

            if primary:
                current_primary = primary
            if secondary:
                current_secondary = secondary
            if tertiary:
                current_tertiary = tertiary
            if not knowledge and not description:
                continue

            dimension = _normalize_dimension(current_primary)
            raw_code = _extract_leading_code(knowledge) or _extract_leading_code(current_tertiary)
            points.append(
                _make_point(
                    source=path.name,
                    stage=sheet_name,
                    primary_dimension=dimension,
                    secondary_dimension=_strip_leading_code(current_secondary),
                    tertiary_ability=_strip_leading_code(current_tertiary),
                    knowledge_point=knowledge or description,
                    description=description,
                    cognitive_level=cognitive_level,
                    suggested_question_types=_split_list(note),
                    source_references=_split_list(source_reference),
                    tags=[sheet_name, cognitive_level],
                    source_sheet=sheet_name,
                    source_row=row_number,
                    sequence=sequence,
                    raw_code=raw_code,
                )
            )
            sequence += 1
    return points


def dedupe_points(points: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    seen: set[str] = set()
    result: List[Dict[str, Any]] = []
    for point in points:
        key = "|".join(
            [
                point["stage"],
                point["primaryDimension"],
                point["secondaryDimension"],
                point["tertiaryAbility"],
                point["knowledgePoint"],
            ]
        )
        if key in seen:
            continue
        seen.add(key)
        result.append(point)
    return result


def upsert_points(points: List[Dict[str, Any]], dry_run: bool = False) -> ImportStats:
    stats = ImportStats(parsed=len(points))
    now = now_iso()
    with sqlite3.connect(str(get_db_path())) as connection:
        connection.row_factory = sqlite3.Row
        for point in points:
            existing = connection.execute(
                "SELECT id FROM knowledge_points WHERE id = ? OR knowledge_code = ?",
                (point["id"], point["knowledgeCode"]),
            ).fetchone()
            if existing:
                stats.updated += 1
            else:
                stats.inserted += 1
            if dry_run:
                continue
            connection.execute(
                """
                INSERT OR REPLACE INTO knowledge_points
                (id, knowledge_code, stage, primary_dimension, secondary_dimension,
                 tertiary_ability, knowledge_point, description, cognitive_level,
                 suggested_question_types_json, source_references_json, tags_json,
                 status, source_sheet, source_row, created_at, updated_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    point["id"],
                    point["knowledgeCode"],
                    point["stage"],
                    point["primaryDimension"],
                    point["secondaryDimension"],
                    point["tertiaryAbility"],
                    point["knowledgePoint"],
                    point["description"],
                    point["cognitiveLevel"],
                    to_json(point["suggestedQuestionTypes"]),
                    to_json(point["sourceReferences"]),
                    to_json(point["tags"]),
                    point["status"],
                    point["sourceSheet"],
                    point["sourceRow"],
                    now,
                    now,
                ),
            )
    return stats


def main() -> int:
    parser = argparse.ArgumentParser(description="Import normalized AI literacy knowledge points into SQLite.")
    parser.add_argument("--root", default="knowledge", help="Folder containing the workbook files.")
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--json-output", default="data/knowledge_points_import.json")
    args = parser.parse_args()

    init_db()
    workbooks = _find_workbooks(Path(args.root))
    points: List[Dict[str, Any]] = []
    if "graph" in workbooks:
        points.extend(parse_graph_workbook(workbooks["graph"]))
    if "design" in workbooks:
        points.extend(parse_design_workbook(workbooks["design"]))
    points = dedupe_points(points)

    output = Path(args.json_output)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(points, ensure_ascii=False, indent=2), encoding="utf-8")
    stats = upsert_points(points, dry_run=args.dry_run)

    by_stage: Dict[str, int] = {}
    by_dimension: Dict[str, int] = {}
    for point in points:
        by_stage[point["stage"]] = by_stage.get(point["stage"], 0) + 1
        by_dimension[point["primaryDimension"]] = by_dimension.get(point["primaryDimension"], 0) + 1

    print(f"dry_run={args.dry_run}")
    print(f"parsed={stats.parsed}")
    print(f"inserted={stats.inserted}")
    print(f"updated={stats.updated}")
    print(f"json_output={output}")
    print("by_stage=", by_stage)
    print("by_dimension=", by_dimension)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
