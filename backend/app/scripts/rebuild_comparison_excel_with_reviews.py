from __future__ import annotations

import json
import re
import shutil
from collections import Counter
from pathlib import Path
from typing import Any, Dict, List, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


INPUT = Path("data/knowledge_point_model_comparison_rows_2026-06-12.json")
OUTPUT = Path("outputs/knowledge_point_model_comparison_reviewed_2026-06-12.xlsx")
DISPLAY_COPY = Path("outputs/知识点总表出题模型对比_结构质量分评_2026-06-12.xlsx")

DIMENSIONS = ["理解AI（Understand）", "应用AI（Apply）", "创造AI（Create）", "AI伦理（Ethics）"]
SOURCES = ["讯飞", "DeepSeek", "Codex"]


def structure_review(row: Dict[str, Any]) -> Tuple[str, str, str]:
    issues = []
    options = row.get("options") or []
    option_ids = {str(option.get("id", "")).strip() for option in options}
    answer = str(row.get("correctAnswer", "")).strip()

    if row.get("questionType") != "单选":
        issues.append("题型不是单选")
    if len(options) != 4:
        issues.append(f"选项数为{len(options)}")
    if answer not in option_ids:
        issues.append("答案未匹配 A/B/C/D")
    if not row.get("knowledgeCode"):
        issues.append("缺少 knowledgeCode")
    if not row.get("primaryDimension") or not row.get("secondaryDimension"):
        issues.append("维度字段不完整")
    if not row.get("scenario") or len(str(row.get("scenario", "")).strip()) < 12:
        issues.append("场景为空或过短")
    if not row.get("question"):
        issues.append("题干为空")
    if not row.get("explanation") or len(str(row.get("explanation", "")).strip()) < 20:
        issues.append("解析为空或过短")

    if not issues:
        return "通过", "结构完整，可进入质量复审。", ""
    if any("答案" in item or "选项数" in item or "题干为空" in item for item in issues):
        return "不通过", "结构存在硬伤，暂不进入质量判断。", "；".join(issues)
    return "需修正", "结构基本可读，但字段或场景需补强。", "；".join(issues)


def quality_review(row: Dict[str, Any], structure_result: str) -> Tuple[str, str, str]:
    if structure_result == "不通过":
        return "C", "暂不建议使用", "结构未通过，先修结构。"

    issues = []
    options = row.get("options") or []
    option_texts = [str(option.get("text", "")).strip() for option in options]
    answer = str(row.get("correctAnswer", "")).strip()
    correct_text = ""
    for option in options:
        if str(option.get("id", "")).strip() == answer:
            correct_text = str(option.get("text", "")).strip()

    combined = " ".join(
        [
            str(row.get("title", "")),
            str(row.get("scenario", "")),
            str(row.get("question", "")),
            str(row.get("explanation", "")),
        ]
    )
    knowledge_point = str(row.get("knowledgePoint", "")).strip()
    tertiary = str(row.get("tertiaryAbility", "")).strip()

    if knowledge_point and knowledge_point not in combined and tertiary and tertiary not in combined:
        issues.append("考点显性不足")
    if re.search(r"下列哪项.*(正确|符合)|哪一种.*(正确|符合)", row.get("question", "")) and not row.get("scenario"):
        issues.append("题目偏定义问法，场景化不足")
    if correct_text and any(text == correct_text for text in option_texts if text != correct_text):
        issues.append("选项存在重复或近重复")
    lengths = [len(text) for text in option_texts if text]
    if lengths and max(lengths) > max(18, min(lengths) * 3):
        issues.append("选项长度差异过大，可能提示答案")
    if "以上" in " ".join(option_texts) or "都不" in " ".join(option_texts):
        issues.append("含兜底选项，需检查区分度")
    explanation = str(row.get("explanation", "")).strip()
    if explanation and not any(marker in explanation for marker in ["因为", "因此", "关键", "错误", "不符合", "体现", "说明"]):
        issues.append("解析可能只给结论，缺少理由")
    if not any(marker in str(row.get("scenario", "")) for marker in ["学生", "同学", "老师", "学校", "班级", "小组", "作业", "课堂", "校园"]):
        issues.append("初中学习场景不够明显")

    if not issues:
        return "A", "可进入人工复审", "质量初筛较好，但仍需人工审题。"
    if len(issues) <= 2:
        return "B", "修改后复审", "；".join(issues)
    return "C", "暂不建议使用", "；".join(issues)


def apply_reviews(rows: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    reviewed = []
    for row in rows:
        row = dict(row)
        structure_result, structure_comment, structure_issues = structure_review(row)
        quality_level, quality_suggestion, quality_issues = quality_review(row, structure_result)
        row["structureResult"] = structure_result
        row["structureComment"] = structure_comment
        row["structureIssues"] = structure_issues
        row["qualityLevel"] = quality_level
        row["qualitySuggestion"] = quality_suggestion
        row["qualityIssues"] = quality_issues
        row["optionsText"] = "\n".join(f"{option.get('id')}. {option.get('text')}" for option in row.get("options", []))
        reviewed.append(row)
    return reviewed


HEADERS = [
    "编号",
    "来源",
    "knowledgeCode",
    "知识条目",
    "学段",
    "一级维度",
    "二级维度",
    "三级能力",
    "四级知识点",
    "题型",
    "题目",
    "场景",
    "选项",
    "正确答案",
    "解析",
    "参考知识点",
    "结构规则结果",
    "结构规则说明",
    "结构问题备注",
    "质量规则等级",
    "质量规则建议",
    "质量问题备注",
]

FIELD_MAP = {
    "编号": "number",
    "来源": "source",
    "knowledgeCode": "knowledgeCode",
    "知识条目": "knowledgeEntry",
    "学段": "stage",
    "一级维度": "primaryDimension",
    "二级维度": "secondaryDimension",
    "三级能力": "tertiaryAbility",
    "四级知识点": "knowledgePoint",
    "题型": "questionType",
    "题目": "question",
    "场景": "scenario",
    "选项": "optionsText",
    "正确答案": "correctAnswer",
    "解析": "explanation",
    "参考知识点": "description",
    "结构规则结果": "structureResult",
    "结构规则说明": "structureComment",
    "结构问题备注": "structureIssues",
    "质量规则等级": "qualityLevel",
    "质量规则建议": "qualitySuggestion",
    "质量问题备注": "qualityIssues",
}


def _fill_for_structure(value: str) -> PatternFill:
    if value == "通过":
        return PatternFill("solid", fgColor="E2F0D9")
    if value == "需修正":
        return PatternFill("solid", fgColor="FFF2CC")
    return PatternFill("solid", fgColor="FCE4D6")


def _fill_for_quality(value: str) -> PatternFill:
    if value == "A":
        return PatternFill("solid", fgColor="DDEBF7")
    if value == "B":
        return PatternFill("solid", fgColor="FFF2CC")
    return PatternFill("solid", fgColor="FCE4D6")


def _style_header(row) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)
    for cell in row:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _style_sheet(ws) -> None:
    thin = Side(style="thin", color="D9E2F3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def build_excel(rows: List[Dict[str, Any]]) -> None:
    wb = Workbook()
    detail = wb.active
    detail.title = "对比明细"
    detail.append(HEADERS)
    _style_header(detail[1])
    for row in rows:
        detail.append([row.get(FIELD_MAP[header], "") for header in HEADERS])

    widths = [15, 12, 22, 28, 10, 20, 28, 28, 22, 14, 42, 42, 52, 12, 54, 42, 16, 32, 36, 16, 18, 40]
    for index, width in enumerate(widths, start=1):
        detail.column_dimensions[get_column_letter(index)].width = width
    for row_index in range(2, detail.max_row + 1):
        detail.row_dimensions[row_index].height = 94
        structure_cell = detail.cell(row=row_index, column=HEADERS.index("结构规则结果") + 1)
        quality_cell = detail.cell(row=row_index, column=HEADERS.index("质量规则等级") + 1)
        structure_cell.fill = _fill_for_structure(str(structure_cell.value))
        quality_cell.fill = _fill_for_quality(str(quality_cell.value))
    _style_sheet(detail)

    summary = wb.create_sheet("汇总")
    summary.append(["来源", "一级维度", "题目数", "结构通过", "结构需修正", "结构不通过", "质量A", "质量B", "质量C"])
    _style_header(summary[1])
    for source in SOURCES:
        for dimension in DIMENSIONS:
            filtered = [row for row in rows if row["source"] == source and row["primaryDimension"] == dimension]
            structure_counts = Counter(row["structureResult"] for row in filtered)
            quality_counts = Counter(row["qualityLevel"] for row in filtered)
            summary.append(
                [
                    source,
                    dimension,
                    len(filtered),
                    structure_counts.get("通过", 0),
                    structure_counts.get("需修正", 0),
                    structure_counts.get("不通过", 0),
                    quality_counts.get("A", 0),
                    quality_counts.get("B", 0),
                    quality_counts.get("C", 0),
                ]
            )
    summary.append([])
    summary.append(["总题数", len(rows)])
    summary.append(["说明", "结构规则只判断字段与格式；质量规则是脚本初筛，不等于人工终审。"])
    for index, width in enumerate([14, 24, 12, 12, 12, 12, 10, 10, 10], start=1):
        summary.column_dimensions[get_column_letter(index)].width = width
    _style_sheet(summary)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)
    shutil.copyfile(OUTPUT, DISPLAY_COPY)

    check = load_workbook(OUTPUT, read_only=True, data_only=True)
    if check["对比明细"].max_row - 1 != len(rows):
        raise RuntimeError("Workbook verification failed")


def main() -> int:
    rows = json.loads(INPUT.read_text(encoding="utf-8"))
    reviewed = apply_reviews(rows)
    INPUT.with_name("knowledge_point_model_comparison_rows_reviewed_2026-06-12.json").write_text(
        json.dumps(reviewed, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    build_excel(reviewed)
    print(f"rows={len(reviewed)}")
    print(f"excel={OUTPUT}")
    print(f"display_copy={DISPLAY_COPY}")
    print("structure_counts=", Counter(row["structureResult"] for row in reviewed))
    print("quality_counts=", Counter(row["qualityLevel"] for row in reviewed))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
